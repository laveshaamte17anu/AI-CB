import os
import re
from typing import List, Dict, Any, Tuple
from pathlib import Path
import logging

logger = logging.getLogger(__name__)

# Supported formats and size limit (25 MB)
SUPPORTED_EXTENSIONS = {".pdf", ".txt", ".docx", ".md"}
MAX_FILE_SIZE = 25 * 1024 * 1024  # 25 MB

def validate_document(filename: str, file_size: int) -> Tuple[bool, str]:
    """Validates document format and size limits."""
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        return False, f"Unsupported file type '{ext}'. Supported formats: PDF, TXT, DOCX, MD."
    if file_size > MAX_FILE_SIZE:
        return False, f"File size exceeds 25 MB limit (file is {file_size / (1024*1024):.1f} MB)."
    return True, ""

def extract_text_from_pdf(file_path: str) -> Tuple[List[Dict[str, Any]], int, str]:
    """
    Extracts text page by page from PDF using pypdf.
    Returns (pages_list, page_count, combined_text).
    """
    from pypdf import PdfReader

    pages = []
    combined_parts = []
    
    try:
        reader = PdfReader(file_path)
        total_pages = len(reader.pages)
        
        has_text = False
        for page_idx, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            text = text.strip()
            if text:
                has_text = True
                pages.append({"page_number": page_idx + 1, "text": text})
                combined_parts.append(f"--- Page {page_idx + 1} ---\n{text}")
            else:
                pages.append({"page_number": page_idx + 1, "text": "[No text extracted on this page]"})

        if not has_text:
            warning = "Notice: This PDF appears to be scanned or contains only images. OCR text extraction requires an image OCR engine."
            combined_parts.append(warning)
            
        return pages, total_pages, "\n\n".join(combined_parts)
    except Exception as e:
        logger.error(f"Error reading PDF {file_path}: {e}")
        raise ValueError(f"Could not read PDF: {str(e)}")

def extract_text_from_docx(file_path: str) -> Tuple[List[Dict[str, Any]], int, str]:
    """Extracts text from DOCX file using python-docx."""
    from docx import Document

    try:
        doc = Document(file_path)
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        
        # Also extract table text
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    paragraphs.append(row_text)
                    
        full_text = "\n\n".join(paragraphs)
        pages = [{"page_number": 1, "text": full_text}]
        return pages, 1, full_text
    except Exception as e:
        logger.error(f"Error reading DOCX {file_path}: {e}")
        raise ValueError(f"Could not read DOCX: {str(e)}")

def extract_text_from_txt(file_path: str) -> Tuple[List[Dict[str, Any]], int, str]:
    """Extracts text from TXT or MD files."""
    try:
        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            text = f.read().strip()
        pages = [{"page_number": 1, "text": text}]
        return pages, 1, text
    except Exception as e:
        logger.error(f"Error reading text file {file_path}: {e}")
        raise ValueError(f"Could not read text file: {str(e)}")

def parse_document(file_path: str, filename: str) -> Tuple[List[Dict[str, Any]], int, str]:
    """Routes document to appropriate parser based on extension."""
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        return extract_text_from_pdf(file_path)
    elif ext == ".docx":
        return extract_text_from_docx(file_path)
    elif ext in {".txt", ".md"}:
        return extract_text_from_txt(file_path)
    else:
        raise ValueError(f"Unsupported file format: {ext}")

def chunk_document_pages(pages: List[Dict[str, Any]], chunk_size_words: int = 350, overlap_words: int = 50) -> List[Dict[str, Any]]:
    """
    Chunks document page text into coherent sections with page references.
    """
    chunks = []
    chunk_index = 0

    for page_data in pages:
        page_num = page_data["page_number"]
        page_text = page_data["text"]
        words = page_text.split()
        
        if not words:
            continue

        if len(words) <= chunk_size_words:
            chunks.append({
                "chunk_index": chunk_index,
                "page_number": page_num,
                "content": page_text
            })
            chunk_index += 1
        else:
            start = 0
            while start < len(words):
                end = start + chunk_size_words
                chunk_slice = words[start:end]
                chunks.append({
                    "chunk_index": chunk_index,
                    "page_number": page_num,
                    "content": " ".join(chunk_slice)
                })
                chunk_index += 1
                start += (chunk_size_words - overlap_words)

    return chunks

def extract_keywords(text: str) -> List[str]:
    """Extracts search keywords from a query string."""
    words = re.findall(r'\b[A-Za-z0-9_]{3,}\b', text.lower())
    # Exclude common stopwords
    stopwords = {
        "the", "and", "for", "with", "that", "this", "from", "what", "how", "why",
        "when", "where", "can", "you", "tell", "about", "explain", "does", "have"
    }
    return [w for w in words if w not in stopwords]
