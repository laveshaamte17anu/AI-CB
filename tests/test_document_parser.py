import sys
import os
import tempfile
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from backend.document_parser import (
    validate_document,
    parse_document,
    chunk_document_pages,
    extract_keywords
)

def test_validate_document():
    valid, _ = validate_document("test.pdf", 1024)
    assert valid is True

    valid_docx, _ = validate_document("report.docx", 2048)
    assert valid_docx is True

    invalid_ext, err = validate_document("malicious.exe", 1024)
    assert invalid_ext is False
    assert "Unsupported file type" in err

    oversized, err_size = validate_document("huge.pdf", 30 * 1024 * 1024)
    assert oversized is False
    assert "exceeds 25 MB" in err_size

def test_chunk_and_extract_text():
    with tempfile.NamedTemporaryFile(mode="w", suffix=".txt", delete=False, encoding="utf-8") as f:
        f.write("Line 1: Introduction to Local AI.\nLine 2: Everything runs locally on device.")
        temp_path = f.name

    try:
        pages, page_count, full_text = parse_document(temp_path, "sample.txt")
        assert page_count == 1
        assert "Introduction to Local AI" in full_text

        chunks = chunk_document_pages(pages, chunk_size_words=5, overlap_words=1)
        assert len(chunks) >= 2
        assert chunks[0]["page_number"] == 1
    finally:
        if os.path.exists(temp_path):
            os.remove(temp_path)

def test_extract_keywords():
    query = "What is the capital of France and how does it work?"
    kw = extract_keywords(query)
    assert "capital" in kw
    assert "france" in kw
    assert "the" not in kw
    assert "what" not in kw

def test_docx_extraction():
    # pyrefly: ignore [missing-import]
    from docx import Document
    doc = Document()
    doc.add_paragraph("Paragraph 1: Testing local document parser.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Header A"
    table.rows[0].cells[1].text = "Header B"

    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
        doc_path = f.name

    try:
        doc.save(doc_path)
        pages, page_count, full_text = parse_document(doc_path, "sample.docx")
        assert page_count == 1
        assert "Paragraph 1: Testing local document parser." in full_text
        assert "Header A" in full_text
    finally:
        if os.path.exists(doc_path):
            os.remove(doc_path)

def test_corrupted_file_handling():
    with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as f:
        f.write(b"not a valid pdf content")
        corrupted_path = f.name

    try:
        # pyrefly: ignore [missing-import]
        import pytest
        with pytest.raises(ValueError) as excinfo:
            parse_document(corrupted_path, "corrupted.pdf")
        assert "Could not read PDF" in str(excinfo.value)
    finally:
        if os.path.exists(corrupted_path):
            os.remove(corrupted_path)

def test_pdf_extraction():
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as f:
        pdf_path = f.name

    c = canvas.Canvas(pdf_path, pagesize=letter)
    c.drawString(100, 700, "Testing PDF text extraction for Local AI Assistant.")
    c.showPage()
    c.save()

    try:
        pages, page_count, full_text = parse_document(pdf_path, "test.pdf")
        assert page_count == 1
        assert "Testing PDF text extraction" in full_text
    finally:
        if os.path.exists(pdf_path):
            os.remove(pdf_path)

